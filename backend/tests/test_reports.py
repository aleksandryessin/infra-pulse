"""R1 management reports without a database: monthly summary, XLSX and journal export.

ТЗ §8 (optional). Counters add up to «выдано», decisions cover R1–R7, an empty month is a
report of zeros, the XLSX renders the same ``MonthlyReport`` (same numbers), dates are
MSK wall time, text never becomes a formula, the journal streams page by page within
93 days, and no session, password, note text or author ID reaches a file. Synthetic data.
"""

import io
import tracemalloc
import zipfile
from datetime import UTC, date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from infra_pulse_backend.api import forecast_fixture
from infra_pulse_backend.api.app import create_app
from infra_pulse_backend.api.auth_deps import current_actor
from infra_pulse_backend.config import Settings
from infra_pulse_backend.ingestion.journal_csv import MSK
from infra_pulse_backend.operations.reports import (
    DECISION_CODES,
    JOURNAL_HEADER,
    MAX_JOURNAL_DAYS,
    AlarmCounts,
    ReportRequestError,
    build_monthly_report,
    card_fact,
    check_journal_period,
    iter_journal,
    journal_xlsx,
    monthly_report_xlsx,
    parse_month,
)
from infra_pulse_core.contracts.auth import Me
from infra_pulse_core.contracts.forecast import FEEDER_TARGET, ForecastDecisionSummary
from infra_pulse_core.contracts.reports import MonthlyReport

GENERATED = datetime(2026, 9, 27, 12, 0, tzinfo=UTC)
ALARMS = AlarmCounts(
    by_type={"Датчик дыма": 7, "Состояние фазы": 3, "тип не указан": 1},
    by_object={"synthetic-object-11": 4, "synthetic-object-99": 9},
)


def fixture_report(month: str = "2026-09", **overrides):
    arguments = {
        "mode": "fixture",
        "month": month,
        "generated_at": GENERATED,
        "data_as_of": forecast_fixture.forecast_state().data_as_of,
        "entries": forecast_fixture.journal_entries(),
        "targets": [FEEDER_TARGET],
        "alarms": ALARMS,
        "object_names": {"synthetic-object-11": "Синтетический объект 11"},
    }
    return build_monthly_report(**(arguments | overrides))


def test_parse_month_and_periods():
    assert parse_month("2026-09") == (date(2026, 9, 1), date(2026, 10, 1))
    assert parse_month("2026-12") == (date(2026, 12, 1), date(2027, 1, 1))
    for bad in ("2026-13", "2026-9", "26-09", "2026-09-01", ""):
        with pytest.raises(ReportRequestError, match="invalid_month"):
            parse_month(bad)
    check_journal_period(date(2026, 7, 1), date(2026, 7, 1) + timedelta(days=MAX_JOURNAL_DAYS - 1))
    with pytest.raises(ReportRequestError, match="journal_period_exceeds_93_days"):
        check_journal_period(date(2026, 7, 1), date(2026, 7, 1) + timedelta(days=MAX_JOURNAL_DAYS))
    with pytest.raises(ReportRequestError, match="issued_range_reversed"):
        check_journal_period(date(2026, 7, 2), date(2026, 7, 1))


def test_monthly_counters_add_up_and_match_the_journal():
    report = fixture_report()
    scored = [
        entry
        for entry in forecast_fixture.journal_entries()
        if entry.card.status == "scored"
        and entry.card.issued_at.astimezone(MSK).strftime("%Y-%m") == "2026-09"
    ]
    (row,) = report.cards
    assert row.target_spec_id == FEEDER_TARGET
    assert row.issued == len(scored)
    assert row.released + row.no_event + row.unknown + row.open == row.issued
    facts = [card_fact(entry) for entry in scored]
    assert (row.released, row.no_event, row.unknown, row.open) == (
        facts.count("released"),
        facts.count("no_event"),
        facts.count("unknown"),
        facts.count("open"),
    )
    assert [item.decision_code for item in report.decisions] == list(DECISION_CODES)
    assert sum(item.count for item in report.decisions) == sum(
        entry.decision is not None for entry in scored
    )
    assert report.basis == "automatic_registered_event"
    assert any("не метрика качества" in note for note in report.notes)
    assert any("fixture" in note for note in report.notes)


def test_top_objects_and_alarm_types():
    report = fixture_report()
    assert len(report.top_objects) <= 10
    keys = [
        (-row.cards, -row.released, -row.source_alarms, row.object_id) for row in report.top_objects
    ]
    assert keys == sorted(keys)
    by_id = {row.object_id: row for row in report.top_objects}
    assert by_id["synthetic-object-11"].source_alarms == 4
    assert by_id["synthetic-object-11"].object_name == "Синтетический объект 11"
    assert [row.sensor_type for row in report.alarms_by_type] == [
        "Датчик дыма",
        "Состояние фазы",
        "тип не указан",
    ]
    # Thirteen objects have cards: an object with source alarms only stays below them.
    assert all(row.cards for row in report.top_objects)
    assert "synthetic-object-99" not in by_id
    few = fixture_report(entries=forecast_fixture.journal_entries()[:2])
    assert [(row.object_id, row.cards, row.source_alarms) for row in few.top_objects[-2:]] == [
        ("synthetic-object-99", 0, 9),
        ("synthetic-object-11", 0, 4),
    ]
    assert all(row.cards for row in few.top_objects[:-2])


def test_empty_month_is_a_report_of_zeros():
    report = fixture_report("2026-02", alarms=None)
    assert [row.model_dump() for row in report.cards] == [
        {
            "target_spec_id": FEEDER_TARGET,
            "issued": 0,
            "released": 0,
            "no_event": 0,
            "unknown": 0,
            "open": 0,
        }
    ]
    assert [row.count for row in report.decisions] == [0] * 7
    assert report.top_objects == [] and report.alarms_by_type == []
    assert report.period_start == date(2026, 2, 1) and report.period_end == date(2026, 3, 1)


def test_unfinished_month_is_noted():
    report = fixture_report()
    assert any("не завершён" in note for note in report.notes)
    closed = fixture_report("2026-08")
    assert not any("не завершён" in note for note in closed.notes)


def sheet_rows(book, title):
    return [list(row) for row in book[title].iter_rows(values_only=True)]


def test_monthly_xlsx_has_the_json_numbers():
    report = fixture_report()
    book = load_workbook(monthly_report_xlsx(report))
    assert book.sheetnames == [
        "Карточки",
        "Решения",
        "Объекты",
        "Тревожные сообщения по типам",
        "Условия",
    ]
    assert all(len(name) <= 31 for name in book.sheetnames)  # Excel limit
    cards = sheet_rows(book, "Карточки")
    assert cards[0] == [
        "Прогноз",
        "Цель",
        "Выдано",
        "Снята по событию",
        "Без события",
        "Неизвестно",
        "Открыто",
    ]
    assert [row[1:] for row in cards[1:]] == [
        [row.target_spec_id, row.issued, row.released, row.no_event, row.unknown, row.open]
        for row in report.cards
    ]
    decisions = sheet_rows(book, "Решения")
    assert decisions[1:] == [[row.decision_code, row.count] for row in report.decisions]
    objects = sheet_rows(book, "Объекты")
    assert objects[1:] == [
        [place, row.object_name or None, row.object_id, row.cards, row.released, row.source_alarms]
        for place, row in enumerate(report.top_objects, start=1)
    ]
    alarms = sheet_rows(book, "Тревожные сообщения по типам")
    assert alarms[0] == ["Тип датчика", "Исходных тревожных сообщений"]
    assert alarms[1:] == [[row.sensor_type, row.source_alarms] for row in report.alarms_by_type]
    conditions = dict((row[0], row[1]) for row in sheet_rows(book, "Условия")[1:7])
    assert conditions["Режим данных"] == "синтетические данные (fixture)"
    assert conditions["Сформирован (МСК)"] == datetime(2026, 9, 27, 15, 0)  # 12:00 UTC
    assert conditions["Месяц"] == "2026-09"


def formula_free(buffer) -> str:
    with zipfile.ZipFile(buffer) as archive:
        xml = "".join(
            archive.read(name).decode()
            for name in archive.namelist()
            if name.startswith("xl/worksheets/")
        )
    assert "<f>" not in xml and "<f " not in xml
    return xml


def decision(reason_text: str) -> ForecastDecisionSummary:
    return ForecastDecisionSummary(
        decision_code="R2",
        reason_code="R2.3",
        reason_text=reason_text,
        dictionary_version="synthetic-technologist-dictionary-v0",
        actor_id="synthetic-secret-actor-77",
        actor_role="dispatcher",
        decided_at=forecast_fixture.FIXTURE_AS_OF,
        revision=1,
        simulated=True,
        verification_methods=["source_records"],
        notified_to='=HYPERLINK("http://example.invalid")',
        notified_at=forecast_fixture.FIXTURE_AS_OF - timedelta(minutes=5),
    )


def journal_entries_with_hostile_text():
    entries = forecast_fixture.journal_entries()
    first = next(entry for entry in entries if entry.card.status == "scored")
    hostile = first.model_copy(
        update={"decision": decision("=1+1 причина\x07 с управляющим символом")}
    )
    return [hostile if entry is first else entry for entry in entries], hostile


def test_journal_xlsx_rows_text_and_dates():
    entries, hostile = journal_entries_with_hostile_text()
    export = journal_xlsx(
        entries,
        mode="fixture",
        issued_from=date(2026, 9, 1),
        issued_to=date(2026, 9, 30),
        generated_at=GENERATED,
        data_as_of=forecast_fixture.forecast_state().data_as_of,
        object_names=lambda ids: {object_id: f"Имя {object_id}" for object_id in ids},
    )
    raw = export.file.read()
    xml = formula_free(io.BytesIO(raw))
    assert "synthetic-secret-actor-77" not in xml  # the author is shown by role only
    book = load_workbook(io.BytesIO(raw))
    rows = sheet_rows(book, "Журнал")
    assert rows[0] == JOURNAL_HEADER
    scored = [entry for entry in entries if entry.card.status == "scored"]
    assert export.rows == len(scored) == len(rows) - 1
    by_card = {row[3]: row for row in rows[1:]}
    row = by_card[hostile.card.id]
    assert row[1] == hostile.card.issued_at.astimezone(MSK).replace(tzinfo=None)
    assert row[14] == "R2.3 =1+1 причина\\x07 с управляющим символом"
    assert row[16] == "диспетчер"
    assert row[17] == '=HYPERLINK("http://example.invalid")'
    assert row[6] == f"Имя {hostile.card.object_id}"
    conditions = dict((line[0], line[1]) for line in sheet_rows(book, "Условия")[1:5])
    assert conditions["Строк"] == export.rows


def test_journal_xlsx_period_filter_and_limit():
    entries = forecast_fixture.journal_entries()
    export = journal_xlsx(
        entries,
        mode="fixture",
        issued_from=date(2026, 9, 20),
        issued_to=date(2026, 9, 21),
        generated_at=GENERATED,
        data_as_of=None,
    )
    expected = [
        entry
        for entry in entries
        if entry.card.status == "scored"
        and date(2026, 9, 20) <= entry.card.issued_at.astimezone(MSK).date() <= date(2026, 9, 21)
    ]
    assert export.rows == len(expected) > 0
    with pytest.raises(ReportRequestError):
        journal_xlsx(
            entries,
            mode="fixture",
            issued_from=date(2026, 1, 1),
            issued_to=date(2026, 9, 30),
            generated_at=GENERATED,
            data_as_of=None,
        )


def test_iter_journal_follows_the_pinned_cursor():
    pages = []

    def fetch(cursor):
        pages.append(cursor)
        return forecast_fixture.forecast_journal(cursor=cursor, limit=3)

    entries = list(iter_journal(fetch))
    assert len(entries) == len(forecast_fixture.journal_entries())
    assert len({entry.journal_position for entry in entries}) == len(entries)
    assert pages[0] is None and len(pages) == -(-len(entries) // 3)


def repeated(entry, count):
    for position in range(count):
        yield entry.model_copy(update={"journal_position": position + 1})


def peak_bytes(count: int) -> int:
    entry = next(e for e in forecast_fixture.journal_entries() if e.card.status == "scored")
    issued = entry.card.issued_at.astimezone(MSK).date()
    tracemalloc.start()
    try:
        export = journal_xlsx(
            repeated(entry, count),
            mode="fixture",
            issued_from=issued,
            issued_to=issued,
            generated_at=GENERATED,
            data_as_of=None,
        )
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert export.rows == count
    return peak


def test_journal_xlsx_memory_does_not_grow_with_rows():
    small, large = peak_bytes(1_000), peak_bytes(6_000)
    assert large < small * 1.5 + 2_000_000


# --- HTTP routes -------------------------------------------------------------------------

NOW = datetime(2026, 9, 27, 12, 0, tzinfo=UTC)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def client_as(mode: str, *roles: str, auth_source: str = "ldap", **settings) -> TestClient:
    app = create_app(Settings(mode=mode, _env_file=None, **settings))
    if roles:
        actor = Me(
            subject_id="synthetic-actor",
            display_name="Синтетический пользователь",
            roles=list(roles),
            auth_source=auth_source,
            session_expires_at=NOW + timedelta(hours=1) if auth_source == "ldap" else None,
        )
        app.dependency_overrides[current_actor] = lambda: actor
    return TestClient(app)


@pytest.mark.parametrize(
    "path",
    [
        "/api/v1/reports/monthly?month=2026-09",
        "/api/v1/reports/monthly.xlsx?month=2026-09",
        "/api/v1/reports/journal.xlsx?issued_from=2026-09-01&issued_to=2026-09-30",
    ],
)
def test_report_permission_is_analyst_or_admin(path):
    for roles, status in (
        (("analyst",), 200),
        (("admin",), 200),
        (("dispatcher",), 403),
    ):
        assert client_as("fixture", *roles).get(path).status_code == status, roles
    integration = client_as("fixture", "integration", auth_source="token").get(path)
    assert integration.status_code == 403


def test_fixture_monthly_xlsx_matches_the_json_route():
    client = client_as("fixture", "analyst")
    report = MonthlyReport.model_validate(
        client.get("/api/v1/reports/monthly", params={"month": "2026-09"}).json()
    )
    response = client.get("/api/v1/reports/monthly.xlsx", params={"month": "2026-09"})
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-disposition"] == (
        'attachment; filename="infrapulse-monthly-2026-09.xlsx"'
    )
    book = load_workbook(io.BytesIO(response.content))
    assert [row[1:] for row in sheet_rows(book, "Карточки")[1:]] == [
        [row.target_spec_id, row.issued, row.released, row.no_event, row.unknown, row.open]
        for row in report.cards
    ]
    assert sheet_rows(book, "Решения")[1:] == [
        [row.decision_code, row.count] for row in report.decisions
    ]
    assert [row[2:] for row in sheet_rows(book, "Объекты")[1:]] == [
        [row.object_id, row.cards, row.released, row.source_alarms] for row in report.top_objects
    ]
    assert sheet_rows(book, "Тревожные сообщения по типам")[1:] == [
        [row.sensor_type, row.source_alarms] for row in report.alarms_by_type
    ]
    formula_free(io.BytesIO(response.content))


def test_fixture_journal_xlsx_route():
    client = client_as("fixture", "admin")
    response = client.get(
        "/api/v1/reports/journal.xlsx",
        params={"issued_from": "2026-09-01", "issued_to": "2026-09-30"},
    )
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-disposition"] == (
        'attachment; filename="infrapulse-journal-2026-09-01-2026-09-30.xlsx"'
    )
    rows = sheet_rows(load_workbook(io.BytesIO(response.content)), "Журнал")
    assert rows[0] == JOURNAL_HEADER
    assert len(rows) - 1 == sum(
        entry.card.status == "scored" for entry in forecast_fixture.journal_entries()
    )
    assert any(row[6] and row[6].startswith("Синтетический объект") for row in rows[1:])
    for params, detail in (
        (
            {"issued_from": "2026-01-01", "issued_to": "2026-09-30"},
            "journal_period_exceeds_93_days",
        ),
        ({"issued_from": "2026-09-30", "issued_to": "2026-09-01"}, "issued_range_reversed"),
    ):
        refused = client.get("/api/v1/reports/journal.xlsx", params=params)
        assert (refused.status_code, refused.json()["detail"]) == (422, detail)


def test_report_routes_answer_503_without_storage():
    # B2 is merged: an unreachable database is a storage outage; without a DSN the
    # forecast read itself is unavailable.
    down = client_as("replay", "analyst", replay_snapshot_id="synthetic", db_dsn="postgresql://x/y")
    no_dsn = client_as("replay", "analyst", replay_snapshot_id="synthetic")
    for path in (
        "/api/v1/reports/monthly?month=2026-09",
        "/api/v1/reports/monthly.xlsx?month=2026-09",
        "/api/v1/reports/journal.xlsx?issued_from=2026-09-01&issued_to=2026-09-30",
    ):
        response = down.get(path)
        assert (response.status_code, response.json()["detail"]) == (
            503,
            "reports_storage_unavailable",
        ), path
        response = no_dsn.get(path)
        assert (response.status_code, response.json()["detail"]) == (
            503,
            "forecast_not_implemented",
        ), path
