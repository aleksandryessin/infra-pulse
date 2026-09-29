"""N1 and R1 on the real published forecast of B2 (no fixture stand-in).

A synthetic received scope with a year of phase history is published by
``operations.forecast_publish.recompute``; one smoke record is a critical source alarm.
The scope is historical: B2's ``ForecastState.published_at`` is the wall-clock time of the
run, far after ``data_as_of``. Notifications must stay on the data timeline (``as_of`` =
the data watermark, contract ``NotificationSummary``), and the monthly report must count
the cards B2 issued. Needs ``INFRA_TEST_RECEIVED_DSN`` (``make check-db``); own schema.
"""

import io
import os
import random
from datetime import datetime, timedelta
from pathlib import Path
from uuid import uuid4

import psycopg
import pytest
from app_role import app_dsn
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from psycopg.conninfo import make_conninfo

from infra_pulse_backend.api.app import create_app
from infra_pulse_backend.config import Settings
from infra_pulse_backend.operations import forecast_publish as fp
from infra_pulse_backend.worker import __main__ as worker_main
from infra_pulse_core.contracts.forecast import ForecastState
from infra_pulse_core.contracts.notifications import NotificationSummary
from infra_pulse_core.contracts.reports import MonthlyReport
from infra_pulse_core.features.phase_feeder_episodes import MSK

MIGRATIONS = Path(__file__).resolve().parents[1] / "migrations"
NAMESPACE = "n1r1-b2"
STREAM = "n1r1-b2-stream"
PHASE = "Состояние фазы"
START = datetime(2024, 1, 1, tzinfo=MSK)
DAYS = 380
END = START + timedelta(days=DAYS) - timedelta(seconds=1)
OBJECTS = [str(930000 + index) for index in range(4)]
NAMES = ["ГРО{n} ПК{p}", "ФВ{n} ПК{p}-{q}", "ФАНС{n} ПК{p}"]
COLUMNS = """namespace_id, snapshot_id, row_uid, channel_id, object_id, sensor_type,
    system_type, value_raw, alarm, event_at, available_at, availability_basis, source_file,
    source_sha256, record_ordinal, event_local_raw, received_position"""


@pytest.fixture
def dsn():
    base = os.environ.get("INFRA_TEST_RECEIVED_DSN")
    if not base:
        pytest.skip("local PostgreSQL integration DSN not provided")
    schema = f"n1r1_{uuid4().hex[:12]}"
    with psycopg.connect(base, autocommit=True) as connection:
        connection.execute(f'CREATE SCHEMA "{schema}"')
    scoped = make_conninfo(base, options=f"-csearch_path={schema}")
    worker_main.apply_migrations(scoped, MIGRATIONS)
    try:
        yield app_dsn(scoped)
    finally:
        with psycopg.connect(base, autocommit=True) as connection:
            connection.execute(f'DROP SCHEMA "{schema}" CASCADE')


def publish_history(dsn: str) -> str:
    """Layout, a year of phase records, one smoke alarm; one recompute. Returns its channel."""
    rng = random.Random(3)
    layout, rows = [], []
    for number, object_id in enumerate(OBJECTS):
        channels = []
        for index, template in enumerate(NAMES):
            channel = f"{object_id}-{index}"
            channels.append(channel)
            name = template.format(n=index + 1, p=10 * index + number, q=10 * index + 5)
            layout.append(
                fp.layout_row(
                    channel,
                    name=name,
                    object_id=object_id,
                    sensor_type=PHASE,
                    system_type="Диспетчерский контроль",
                    tag=f"SYN-{index}",
                    reference_version="synthetic-ref-v1",
                )
            )
        for day in range(DAYS):
            base = START + timedelta(days=day)
            for channel in channels:
                rows.append((channel, object_id, PHASE, base + timedelta(hours=6), "Есть питание"))
                if rng.random() < 0.04 + 0.03 * number:
                    start = base + timedelta(hours=10)
                    rows.append((channel, object_id, PHASE, start, "Неисправен"))
                    rows.append(
                        (channel, object_id, PHASE, start + timedelta(minutes=2), "Обесточен")
                    )
                    rows.append((channel, object_id, PHASE, start + timedelta(hours=2), "Норма"))
    smoke = f"{OBJECTS[0]}-smoke"
    rows.append((smoke, OBJECTS[0], "Датчик дыма", END - timedelta(hours=3), "Обнаружен дым"))
    with psycopg.connect(dsn) as connection:
        connection.execute(
            """INSERT INTO dispatch_replay_snapshots (namespace_id, snapshot_id,
                 manifest_sha256, window_start, window_end, row_count, alarm_count, scope_kind)
               VALUES (%s, %s, NULL, %s, %s, %s, 0, 'received')""",
            (NAMESPACE, STREAM, START, END, len(rows)),
        )
        fp.pg.upsert_layout(connection, layout)
        with connection.cursor() as cursor:
            with cursor.copy(f"COPY dispatch_observations ({COLUMNS}) FROM STDIN") as copy:
                for position, (channel, object_id, sensor, at, value) in enumerate(rows, 1):
                    copy.write_row(
                        (
                            NAMESPACE,
                            STREAM,
                            f"row-{position}",
                            channel,
                            object_id,
                            sensor,
                            "Диспетчерский контроль",
                            value,
                            value == "Неисправен",
                            at,
                            at,
                            "observed",
                            "synthetic.csv",
                            "0" * 64,
                            position,
                            at.strftime("%d.%m.%Y %H:%M:%S"),
                            position,
                        )
                    )
    with psycopg.connect(dsn, autocommit=True) as connection:
        fp.recompute(
            connection,
            data_as_of=END,
            namespace_id=NAMESPACE,
            snapshot_id=STREAM,
            list_from=(START + timedelta(days=366)).date(),
        )
    return smoke


def test_notifications_and_reports_follow_the_b2_publication(dsn):
    smoke = publish_history(dsn)
    settings = Settings(
        mode="received",
        db_dsn=dsn,
        received_namespace=NAMESPACE,
        received_stream_id=STREAM,
        _env_file=None,
    )
    client = TestClient(create_app(settings))
    state = ForecastState.model_validate(client.get("/api/v1/forecast-state").json())
    assert state.generation == 1 and state.data_as_of == END
    # B2 publishes on the wall clock; the data of this scope is historical.
    assert state.published_at > state.data_as_of + timedelta(days=365)
    with psycopg.connect(dsn) as connection:
        cards = connection.execute(
            """SELECT card_id, published_at, issued_at FROM forecast_cards
               WHERE namespace_id = %s AND snapshot_id = %s AND status = 'scored'""",
            (NAMESPACE, STREAM),
        ).fetchall()
    assert cards, "a year of history gives scored cards"

    since = END - timedelta(days=3)
    response = client.get("/api/v1/notifications", params={"since": since.isoformat()})
    assert response.status_code == 200, response.text
    summary = NotificationSummary.model_validate(response.json())
    assert summary.as_of == END  # the data watermark, not the wall-clock publication
    expected = {card_id for card_id, published, _ in cards if since < published <= END}
    assert expected and summary.new_cards == len(expected)
    assert {item.ref_id for item in summary.items if item.kind == "new_card"} == expected
    assert summary.critical_alarms == 1
    [alarm] = [item for item in summary.items if item.kind == "critical_alarm"]
    assert alarm.at == END - timedelta(hours=3) and alarm.object_id == OBJECTS[0]
    assert smoke  # the record is the one inserted with the history
    later = client.get(
        "/api/v1/notifications", params={"since": (END + timedelta(seconds=1)).isoformat()}
    )
    assert (later.status_code, later.json()["detail"]) == (422, "since_after_as_of")

    month = f"{END.astimezone(MSK):%Y-%m}"
    report = MonthlyReport.model_validate(
        client.get("/api/v1/reports/monthly", params={"month": month}).json()
    )
    assert report.data_as_of == END
    issued = sum(row.issued for row in report.cards)
    in_month = sum(1 for _, _, issued_at in cards if f"{issued_at.astimezone(MSK):%Y-%m}" == month)
    assert issued == in_month and issued > 0
    for row in report.cards:
        assert row.released + row.no_event + row.unknown + row.open == row.issued
    workbook = client.get("/api/v1/reports/monthly.xlsx", params={"month": month})
    assert workbook.status_code == 200 and workbook.content.startswith(b"PK")
    journal = client.get(
        "/api/v1/reports/journal.xlsx",
        params={
            "issued_from": f"{END - timedelta(days=13):%Y-%m-%d}",
            "issued_to": f"{END:%Y-%m-%d}",
        },
    )
    assert journal.status_code == 200, journal.text
    sheet = load_workbook(io.BytesIO(journal.content), read_only=True).worksheets[0]
    assert next(sheet.iter_rows(max_row=1, values_only=True))  # a header row
