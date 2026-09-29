"""PostgreSQL path of R1 management reports: source alarms of a month and object names.

Needs ``INFRA_TEST_RECEIVED_DSN`` on a disposable database (``make check-db``); every test
works in its own schema. Observation rows are synthetic; the published forecast state
and journal are the synthetic fixture standing in for B2 (``forecast_db`` is patched).
Checks: alarms by type and object inside the month and not later than the stand clock,
``alarm=false`` and other months excluded, an empty month of zeros, names from the
current objects reference, and the journal XLSX with names resolved per page.
"""

import io
import os
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import psycopg
import pytest
from app_role import app_dsn
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from psycopg.conninfo import make_conninfo

from infra_pulse_backend.api import forecast_db, forecast_fixture, reports_db
from infra_pulse_backend.api.app import create_app
from infra_pulse_backend.config import Settings
from infra_pulse_backend.ingestion.journal_csv import MSK
from infra_pulse_backend.operations.reports import build_monthly_report
from infra_pulse_backend.storage import reports_pg
from infra_pulse_backend.worker import __main__ as worker_main
from infra_pulse_core.contracts.forecast import FEEDER_TARGET
from infra_pulse_core.contracts.reports import MonthlyReport

MIGRATIONS = Path(__file__).resolve().parents[1] / "migrations"
NAMESPACE = "r1-test"
SCOPE = "r1-scope"
STATE = forecast_fixture.forecast_state()
DATA_AS_OF = STATE.data_as_of
NOW = datetime(2026, 9, 27, 12, 0, tzinfo=UTC)
SEPTEMBER = datetime(2026, 9, 1, tzinfo=MSK)
_COLUMNS = (
    "namespace_id, snapshot_id, row_uid, channel_id, object_id, sensor_type, system_type, "
    "value_raw, alarm, event_at, available_at, availability_basis, source_file, "
    "source_sha256, record_ordinal, event_local_raw"
)


@pytest.fixture
def dsn():
    base = os.environ.get("INFRA_TEST_RECEIVED_DSN")
    if not base:
        pytest.skip("local PostgreSQL integration DSN not provided")
    schema = f"r1_{uuid4().hex[:12]}"
    with psycopg.connect(base, autocommit=True) as connection:
        connection.execute(f'CREATE SCHEMA "{schema}"')
    scoped = make_conninfo(base, options=f"-csearch_path={schema}")
    worker_main.apply_migrations(scoped, MIGRATIONS)
    try:
        yield app_dsn(scoped)
    finally:
        with psycopg.connect(base, autocommit=True) as connection:
            connection.execute(f'DROP SCHEMA "{schema}" CASCADE')


@pytest.fixture(autouse=True)
def published(monkeypatch):
    def forecast_journal(settings, *, as_of, records_as_of=None, **filters):
        return forecast_fixture.forecast_journal(**filters)

    monkeypatch.setattr(forecast_db, "forecast_state", lambda settings: STATE)
    monkeypatch.setattr(forecast_db, "forecast_journal", forecast_journal)


def replay_settings(dsn: str) -> Settings:
    return Settings(
        mode="replay",
        db_dsn=dsn,
        replay_namespace=NAMESPACE,
        replay_snapshot_id=SCOPE,
        _env_file=None,
    )


def load_scope(dsn: str, rows: list[tuple]) -> None:
    """rows: (sensor_type, object_id, alarm, event_at, available_at or None)."""
    with psycopg.connect(dsn) as connection:
        connection.execute(
            """INSERT INTO dispatch_replay_snapshots
               (namespace_id, snapshot_id, manifest_sha256, window_start, window_end,
                row_count, alarm_count, scope_kind)
               VALUES (%s, %s, NULL, %s, %s, 0, 0, 'replay')""",
            (NAMESPACE, SCOPE, SEPTEMBER - timedelta(days=40), DATA_AS_OF + timedelta(days=9)),
        )
        connection.cursor().executemany(
            f"""INSERT INTO dispatch_observations ({_COLUMNS})
                VALUES (%s, %s, %s, %s, %s, %s, 'synthetic', 'Не замкнут', %s, %s, %s,
                        'simulated', 'synthetic.csv', %s, %s, 'synthetic')""",
            [
                (
                    NAMESPACE,
                    SCOPE,
                    f"synthetic-r1-{index:04d}",
                    f"synthetic-channel-{index:04d}",
                    object_id,
                    sensor_type,
                    alarm,
                    event_at,
                    available_at or event_at + timedelta(seconds=5),
                    "0" * 64,
                    index,
                )
                for index, (sensor_type, object_id, alarm, event_at, available_at) in enumerate(
                    rows, start=1
                )
            ],
        )


def load_object_names(dsn: str, names: dict[str, str]) -> None:
    with psycopg.connect(dsn) as connection:
        connection.execute(
            """INSERT INTO import_files (import_id, format, file_name, sha256, size_bytes,
                 stored_name, uploaded_by, status, finished_at)
               VALUES ('synthetic-objects', 'reference_objects_csv', 'objects.csv', %s, 1,
                 'objects.csv', 'synthetic', 'published', clock_timestamp())""",
            ("b" * 64,),
        )
        connection.execute(
            """INSERT INTO ref_versions (version_id, kind, sha256, import_id, row_count)
               VALUES ('synthetic-objects-v1', 'objects', %s, 'synthetic-objects', %s)""",
            ("b" * 64, len(names)),
        )
        for line, (object_id, name) in enumerate(names.items(), start=1):
            connection.execute(
                """INSERT INTO ref_objects (version_id, line_no, object_id, level_raw,
                     parent_raw, object_kind, name)
                   VALUES ('synthetic-objects-v1', %s, %s, '1', '', 'synthetic', %s)""",
                (line, object_id, name),
            )


def september_rows() -> list[tuple]:
    inside = DATA_AS_OF - timedelta(days=2)
    return [
        ("КД АВ", "synthetic-object-11", True, inside, None),
        ("КД АВ", "synthetic-object-11", True, inside + timedelta(hours=1), None),
        ("Датчик дыма", "synthetic-object-12", True, inside, None),
        (None, None, True, inside, None),  # no type and no object: counted by type only
        ("КД АВ", "synthetic-object-11", False, inside, None),  # alarm=false
        ("КД АВ", "synthetic-object-11", True, SEPTEMBER - timedelta(minutes=1), None),  # August
        ("КД АВ", "synthetic-object-11", True, DATA_AS_OF + timedelta(hours=1), None),  # future
        # A past event that becomes available after the stand clock (replay availability).
        ("КД АВ", "synthetic-object-11", True, inside, DATA_AS_OF + timedelta(hours=2)),
        ("КД АВ", "synthetic-object-11", True, SEPTEMBER, None),  # first second of the month
    ]


def test_source_alarm_counts_follow_month_and_stand_clock(dsn):
    load_scope(dsn, september_rows())
    counts = reports_pg.read_source_alarm_counts(
        dsn,
        namespace_id=NAMESPACE,
        snapshot_id=SCOPE,
        mode="replay",
        start=SEPTEMBER,
        end=datetime(2026, 10, 1, tzinfo=MSK),
        data_as_of=DATA_AS_OF,
        available_as_of=DATA_AS_OF,
    )
    assert counts.by_type == {"КД АВ": 3, "Датчик дыма": 1, "тип не указан": 1}
    assert counts.by_object == {"synthetic-object-11": 3, "synthetic-object-12": 1}
    with pytest.raises(LookupError):
        reports_pg.read_source_alarm_counts(
            dsn,
            namespace_id=NAMESPACE,
            snapshot_id="missing-scope",
            mode="replay",
            start=SEPTEMBER,
            end=datetime(2026, 10, 1, tzinfo=MSK),
            data_as_of=DATA_AS_OF,
            available_as_of=DATA_AS_OF,
        )


def test_monthly_report_from_journal_and_scope(dsn):
    load_scope(dsn, september_rows())
    load_object_names(dsn, {"synthetic-object-11": "Синтетическая станция 11"})
    report = reports_db.monthly_report(replay_settings(dsn), "2026-09", now=NOW)
    expected = build_monthly_report(
        mode="replay",
        month="2026-09",
        generated_at=NOW,
        data_as_of=DATA_AS_OF,
        entries=forecast_fixture.journal_entries(),
        targets=[FEEDER_TARGET],
    )
    assert report.mode == "replay" and report.data_as_of == DATA_AS_OF
    assert report.cards == expected.cards and report.decisions == expected.decisions
    assert [(row.sensor_type, row.source_alarms) for row in report.alarms_by_type] == [
        ("КД АВ", 3),
        ("Датчик дыма", 1),
        ("тип не указан", 1),
    ]
    by_id = {row.object_id: row for row in report.top_objects}
    assert by_id["synthetic-object-11"].source_alarms == 3
    assert by_id["synthetic-object-11"].object_name == "Синтетическая станция 11"
    assert all(
        row.object_name is None
        for row in report.top_objects
        if row.object_id != "synthetic-object-11"
    )


def test_empty_month_is_zeros(dsn):
    load_scope(dsn, september_rows())
    report = reports_db.monthly_report(replay_settings(dsn), "2026-02", now=NOW)
    assert [(row.target_spec_id, row.issued, row.open) for row in report.cards] == [
        (FEEDER_TARGET, 0, 0)
    ]
    assert [row.count for row in report.decisions] == [0] * 7
    assert report.top_objects == [] and report.alarms_by_type == []


def test_journal_export_resolves_names_from_the_reference(dsn):
    load_scope(dsn, [])
    load_object_names(dsn, {"synthetic-object-19": "Синтетическая станция 19"})
    export = reports_db.journal_export(
        replay_settings(dsn), issued_from=date(2026, 9, 1), issued_to=date(2026, 9, 30), now=NOW
    )
    rows = list(load_workbook(export.file)["Журнал"].iter_rows(values_only=True))
    scored = [
        entry for entry in forecast_fixture.journal_entries() if entry.card.status == "scored"
    ]
    assert export.rows == len(scored) == len(rows) - 1
    names = {row[7]: row[6] for row in rows[1:]}
    assert names["synthetic-object-19"] == "Синтетическая станция 19"
    assert names["synthetic-object-11"] is None


def test_routes_json_and_xlsx_have_the_same_numbers(dsn):
    load_scope(dsn, september_rows())
    load_object_names(dsn, {"synthetic-object-11": "Синтетическая станция 11"})
    client = TestClient(create_app(replay_settings(dsn)))
    response = client.get("/api/v1/reports/monthly", params={"month": "2026-09"})
    assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    report = MonthlyReport.model_validate(response.json())
    export = client.get("/api/v1/reports/monthly.xlsx", params={"month": "2026-09"})
    assert export.status_code == 200, export.text
    book = load_workbook(io.BytesIO(export.content))

    def rows(title):
        return [list(line) for line in book[title].iter_rows(values_only=True)][1:]

    assert [line[2:] for line in rows("Карточки")] == [
        [row.issued, row.released, row.no_event, row.unknown, row.open] for row in report.cards
    ]
    assert rows("Решения") == [[row.decision_code, row.count] for row in report.decisions]
    assert [line[1:] for line in rows("Объекты")] == [
        [row.object_name, row.object_id, row.cards, row.released, row.source_alarms]
        for row in report.top_objects
    ]
    assert rows("Тревожные сообщения по типам") == [
        [row.sensor_type, row.source_alarms] for row in report.alarms_by_type
    ]
    journal = client.get(
        "/api/v1/reports/journal.xlsx",
        params={"issued_from": "2026-09-01", "issued_to": "2026-09-30"},
    )
    assert journal.status_code == 200 and journal.headers["cache-control"] == "no-store"
