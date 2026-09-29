"""XLSX rows as text cells, for the same parsers as CSV (G1, ТЗ §7 «импорт CSV/XLSX»).

Only the worker reads workbooks; HTTP stores the uploaded bytes unparsed. The first
worksheet is read with ``openpyxl`` in read-only mode, row by row, cached formula
values only (``data_only``); pandas is not used. Each cell becomes the text a CSV
export of the same column would carry:

* a text cell is kept verbatim (``"01.01.1970 03:00:00"``, ``"25,00"``, leading zeros);
* an integer or an integral number is written without a fraction: ``56682``, never
  ``56682.0`` (identifiers stay text, as in CSV); other numbers use the shortest
  round-trip decimal with a dot (``25.4``);
* ``TRUE``/``FALSE`` become ``true``/``false``;
* an Excel date whose format shows no time becomes ``YYYY-MM-DD``, a time
  ``HH:MM:SS[.ffffff]`` (the
  ``дата``/``время`` columns of the organizers' journal); a date with time becomes
  ``DD.MM.YYYY HH:MM:SS[.ffffff]``, the source system's own text for timestamps (so the
  device placeholder ``01.01.1970 03:00:00`` stays recognisable, and ТЗ Appendix 1
  «Дата записи» reads the same as in CSV). Excel's display format is not copied:
  the stored value is.

Values that Excel itself changed when a CSV was opened in it (``1,5`` → 1.5, dates
re-typed) cannot be recovered here; text-formatted columns avoid that.

A workbook is a ZIP archive: before parsing, its declared uncompressed size and the
sheet's row count are bounded, so a small file cannot expand past the worker memory.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time

from infra_pulse_backend.ingestion.csv_source import FileRejected

# A 50 MB CSV holds ~0.9 M journal rows; XLSX of the same rows is far smaller, so the
# row bound (not the byte limit) is what keeps the parsed file within the worker.
MAX_XLSX_ROWS = 1_000_000
MAX_XLSX_UNCOMPRESSED = 512 * 1024 * 1024


_LITERALS = re.compile(r'"[^"]*"|\\.|\[\$[^\]]*\]')


def _has_time(number_format: str | None) -> bool:
    """True when the Excel number format shows a time of day (``h``/``s`` tokens)."""
    tokens = _LITERALS.sub("", (number_format or "").lower())
    return "h" in tokens or "s" in tokens


def _fraction(microsecond: int) -> str:
    return f".{microsecond:06d}".rstrip("0") if microsecond else ""


def cell_text(value: object, number_format: str | None = None) -> str:
    """Text of one cell value as openpyxl returns it (see the module rules).

    ``number_format`` only separates a date at midnight from a date without time.
    """
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.time() == time(0) and not _has_time(number_format):
            return f"{value:%Y-%m-%d}"
        return f"{value:%d.%m.%Y %H:%M:%S}{_fraction(value.microsecond)}"
    if isinstance(value, date):
        return f"{value:%Y-%m-%d}"
    if isinstance(value, time):
        return f"{value:%H:%M:%S}{_fraction(value.microsecond)}"
    return str(value)


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            expanded = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as error:
        raise FileRejected("unknown_format", "not a readable XLSX workbook") from error
    if "xl/workbook.xml" not in names:
        raise FileRejected("unknown_format", "ZIP archive is not an XLSX workbook")
    if expanded > MAX_XLSX_UNCOMPRESSED:
        raise FileRejected("file_too_large", "workbook expands past the XLSX limit")


def xlsx_rows(data: bytes) -> Iterator[tuple[int, list[str]]]:
    """Yield ``(row_number, cells)`` of the first worksheet; rows keep sheet numbering."""
    _check_archive(data)
    import openpyxl  # worker-only import; the HTTP path never parses uploads

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - any openpyxl failure is a bad file
        raise FileRejected("unknown_format", "workbook cannot be opened") from error
    try:
        if not workbook.worksheets:
            raise FileRejected("bad_header", "workbook has no worksheet")
        sheet = workbook.worksheets[0]
        declared = sheet.max_row
        if declared is not None and declared > MAX_XLSX_ROWS + 1:
            raise FileRejected("file_too_large", "worksheet has too many rows")
        rows = sheet.iter_rows()
        for number, row in enumerate(rows, start=1):
            if number > MAX_XLSX_ROWS + 1:
                raise FileRejected("file_too_large", "worksheet has too many rows")
            yield (
                number,
                [cell_text(cell.value, getattr(cell, "number_format", None)) for cell in row],
            )
    finally:
        workbook.close()


__all__ = ["MAX_XLSX_ROWS", "MAX_XLSX_UNCOMPRESSED", "cell_text", "xlsx_rows"]
